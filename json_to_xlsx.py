"""
Convert parsed FedEx rates JSON into a matrix-style XLSX workbook.

Unified export/import table with shared Transport cost columns.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rate_metadata import (
    chart_for_block,
    country_from_source,
    direction_label,
    is_postal_zone_region,
    package_cost_key,
    parse_postal_entry,
    postal_zone_label,
    service_rate_code,
    shipment_fields_for_region,
    shorten_service_name,
    transport_cost_title,
    zone_letter_from_postal_region,
    zone_region_label,
)

SHIPMENT_HEADERS = [
    "Origin Country",
    "Origin Country Region",
    "Origin Postal Code Zone",
    "Destination Country",
    "Destination Country Region",
    "Destination Postal Code Zone",
    "Service",
]

COST_KEY_ORDER = [
    "envelope",
    "pak",
    "package",
    "package_economy",
    "package_regional_economy",
    "package_freight",
]
WEIGHT_RANGE_RE = re.compile(
    r"^\s*([\d,]+(?:\.\d+)?)\s*-\s*([\d,]+(?:\.\d+)?)\s*$"
)


def _parse_number(text: str) -> float:
    text = text.strip()
    if "." in text and "," in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) == 2 and len(parts[1]) <= 2:
            text = text.replace(",", ".")
        else:
            text = text.replace(",", "")
    return float(text)


def format_weight_label(
    entry: dict[str, Any],
    pricing_model: str,
    *,
    open_ended_replacement: dict[str, str] | None = None,
) -> str:
    rate_type = (entry.get("rate_type") or "").lower()
    weight = (entry.get("weight") or "").strip()

    if rate_type == "envelope" and not weight:
        return "Flat"

    if weight and WEIGHT_RANGE_RE.match(weight):
        if open_ended_replacement and weight in open_ended_replacement:
            return open_ended_replacement[weight]
        parts = re.split(r"\s*-\s*", weight)
        lower = _parse_number(parts[0])
        upper = _parse_number(parts[1])
        if upper >= 99999:
            if lower == int(lower):
                return f">={int(lower)}"
            return f">={lower}"
        if upper == int(upper):
            return str(int(upper))
        return str(upper)

    if weight:
        value = _parse_number(weight)
        if value == int(value):
            return str(int(value))
        return str(value)

    if pricing_model in ("per_pound", "per_kilogram"):
        return "Per kg"
    return ""


# When a per-kg block is parsed alone (no companion bracket rows), use these
# open-ended labels so Flat vs p/unit columns stay consistent with net rates.
DEFAULT_OPEN_ENDED_LABEL: dict[str, str] = {
    "package": ">70.5",
    "package_regional_economy": ">70",
}


def open_ended_weight_labels(
    entries: list[dict[str, Any]],
    *,
    cost_key: str = "",
) -> dict[str, str]:
    open_ranges: list[str] = []
    bounded_ranges: list[tuple[str, float, float]] = []

    for entry in entries:
        weight = (entry.get("weight") or "").strip()
        if not weight or not WEIGHT_RANGE_RE.match(weight):
            continue
        parts = re.split(r"\s*-\s*", weight)
        lower = _parse_number(parts[0])
        upper = _parse_number(parts[1])
        if upper >= 99999:
            open_ranges.append(weight)
        else:
            bounded_ranges.append((weight, lower, upper))

    replacements: dict[str, str] = {}

    # Package Economy: dual per-kg bands → <9999 / >=9999
    has_9999_band = any(abs(upper - 9999.0) < 0.01 for _, _, upper in bounded_ranges)
    has_open_after_9999 = any(
        _parse_number(re.split(r"\s*-\s*", w)[0]) >= 9999.0 for w in open_ranges
    )
    if cost_key == "package_economy" or (has_9999_band and has_open_after_9999):
        for weight, _lower, upper in bounded_ranges:
            if abs(upper - 9999.0) < 0.01:
                replacements[weight] = "<9999"
        for open_weight in open_ranges:
            replacements[open_weight] = ">=9999"
        if replacements:
            return replacements

    if not open_ranges:
        return replacements

    if len(open_ranges) == 1 and not bounded_ranges:
        default_label = DEFAULT_OPEN_ENDED_LABEL.get(cost_key)
        if default_label:
            replacements[open_ranges[0]] = default_label
            return replacements
        parts = re.split(r"\s*-\s*", open_ranges[0])
        lower = _parse_number(parts[0])
        if lower == int(lower):
            replacements[open_ranges[0]] = f">={int(lower)}"
        else:
            replacements[open_ranges[0]] = f">={lower}"
        return replacements

    numeric_weights: list[float] = []
    for entry in entries:
        weight = (entry.get("weight") or "").strip()
        if not weight:
            continue
        if WEIGHT_RANGE_RE.match(weight):
            parts = re.split(r"\s*-\s*", weight)
            upper = _parse_number(parts[1])
            if upper < 99999:
                numeric_weights.append(upper)
        else:
            try:
                numeric_weights.append(_parse_number(weight))
            except ValueError:
                pass

    last_numeric = max(numeric_weights, default=None)

    for open_weight in open_ranges:
        parts = re.split(r"\s*-\s*", open_weight)
        lower = _parse_number(parts[0])
        if last_numeric is not None:
            threshold = last_numeric
        else:
            threshold = lower - 0.5 if lower > 0 else lower
        if threshold == int(threshold):
            replacements[open_weight] = f">{int(threshold)}"
        else:
            replacements[open_weight] = f">{threshold}"

    return replacements


def weight_sort_key(label: str) -> tuple[int, float, str]:
    lower = label.lower()
    if lower == "flat":
        return (0, 0.0, label)
    if lower == "per kg":
        return (4, 99999.0, label)
    match = re.match(r"<(\d+(?:\.\d+)?)", label)
    if match:
        return (3, float(match.group(1)) - 0.01, label)
    match = re.match(r">=(\d+(?:\.\d+)?)", label)
    if match:
        return (3, float(match.group(1)), label)
    match = re.match(r">(\d+(?:\.\d+)?)", label)
    if match:
        return (3, float(match.group(1)), label)
    try:
        return (1, float(label), label)
    except ValueError:
        return (2, 0.0, label)


def pricing_label_for_group(
    weight_labels: list[str], pricing_model: str
) -> str:
    if pricing_model in ("per_pound", "per_kilogram"):
        return "Rate by Weight/chargeable kg"
    has_brackets = any(label.lower() not in ("flat", "per kg") for label in weight_labels)
    if has_brackets:
        return "Rate by Weight/chargeable kg"
    return "Rate by: per shipment"


def _is_per_unit_weight_label(cost_key: str, weight_label: str) -> bool:
    """Weight brackets are Flat; only designated open-ended bands are p/unit."""
    label = weight_label.strip()
    lower = label.lower()
    if lower == "flat":
        return False
    if lower == "per kg":
        return True

    if cost_key == "package_economy":
        return label in ("<9999", ">=9999")
    if cost_key == "package":
        return label == ">70.5" or lower.startswith(">")
    if cost_key == "package_regional_economy":
        return label == ">70" or lower.startswith(">")
    if cost_key in ("pak", "envelope"):
        return False
    # Freight (and unknown) rates are per-unit / per kg.
    return True


def value_header_for_group(weight_labels: list[str], cost_key: str = "") -> str:
    if not weight_labels:
        return "Flat"
    if all(not _is_per_unit_weight_label(cost_key, label) for label in weight_labels):
        return "Flat"
    if all(_is_per_unit_weight_label(cost_key, label) for label in weight_labels):
        return "p/unit"
    return "p/unit"


def value_headers_for_labels(
    weight_labels: list[str], cost_key: str = ""
) -> list[str]:
    return [
        "p/unit" if _is_per_unit_weight_label(cost_key, label) else "Flat"
        for label in weight_labels
    ]


def resolve_cost_key(block: dict[str, Any], rate_type: str) -> str:
    if rate_type == "envelope":
        return "envelope"
    if rate_type == "pak":
        return "pak"
    return package_cost_key(block.get("service") or "")


@dataclass
class CostColumnGroup:
    cost_key: str
    pricing_model: str
    title: str
    pricing_label: str
    value_header: str
    currency: str
    weight_labels: list[str]
    prices: dict[str, dict[str, float]] = field(default_factory=dict)

    @property
    def width(self) -> int:
        return 1 + len(self.weight_labels)

    def price_for(self, row_key: str, weight: str) -> float | None:
        return self.prices.get(row_key, {}).get(weight)


@dataclass
class MatrixRow:
    origin_country: str
    origin_country_region: str
    origin_postal_zone: str
    destination_country: str
    destination_country_region: str
    destination_postal_zone: str
    service: str
    currency: str
    row_key: str


def row_key_for(service: str, shipment: dict[str, str]) -> str:
    return "|".join(
        [
            service,
            shipment.get("origin_country_region", ""),
            shipment.get("origin_postal_zone", ""),
            shipment.get("destination_country_region", ""),
            shipment.get("destination_postal_zone", ""),
        ]
    )


def _zone_letter_sort_key(letter: str) -> tuple[int, str]:
    return (len(letter), letter)


def _region_zone_letter(region: str) -> str:
    if is_postal_zone_region(region):
        return zone_letter_from_postal_region(region) or ""
    match = re.search(r" Zone ([A-Z0-9]+)$", region)
    return match.group(1) if match else ""


def _row_primary_region(row: MatrixRow) -> str:
    if row.origin_country:
        return row.destination_country_region or row.destination_postal_zone
    return row.origin_country_region or row.origin_postal_zone


def _row_is_postal(row: MatrixRow) -> bool:
    if row.origin_country:
        return bool(row.destination_postal_zone) and not row.destination_country_region
    return bool(row.origin_postal_zone) and not row.origin_country_region


def _row_sort_key(row: MatrixRow) -> tuple[Any, ...]:
    region = _row_primary_region(row)
    return (
        row.service,
        0 if row.origin_country else 1,
        _zone_letter_sort_key(_region_zone_letter(region)),
        1 if _row_is_postal(row) else 0,
        region,
    )


def _row_direction(row: MatrixRow) -> str:
    return "Export" if row.origin_country else "Import"


def asymmetric_zone_row_keys(rows: list[MatrixRow]) -> set[str]:
    """Row keys whose zone letter exists on only one of import/export for a service."""
    zones_by_service: dict[str, dict[str, set[str]]] = defaultdict(
        lambda: {"Export": set(), "Import": set()}
    )
    row_zones: list[tuple[str, str, str, str]] = []

    for row in rows:
        letter = _region_zone_letter(_row_primary_region(row))
        if not letter:
            continue
        direction = _row_direction(row)
        zones_by_service[row.service][direction].add(letter)
        row_zones.append((row.row_key, row.service, direction, letter))

    asymmetric: set[str] = set()
    for row_key, service, direction, letter in row_zones:
        export_zones = zones_by_service[service]["Export"]
        import_zones = zones_by_service[service]["Import"]
        if not export_zones or not import_zones:
            continue
        other = "Import" if direction == "Export" else "Export"
        if letter not in zones_by_service[service][other]:
            asymmetric.add(row_key)
    return asymmetric


def build_regions_for_block(
    block: dict[str, Any],
    zone_charts: list[dict[str, Any]],
) -> list[str]:
    chart = chart_for_block(block, zone_charts)
    if not chart:
        zones = block.get("rates", {}).get("zones") or []
        return [zone_region_label(block, zone) for zone in zones]

    code = service_rate_code(block.get("service") or "")
    direction = direction_label(block)
    regions: list[str] = []
    seen: set[str] = set()

    postal_by_zone: dict[str, list[str]] = defaultdict(list)
    postal_seen: dict[str, set[str]] = defaultdict(set)
    for entry in chart.get("countries") or []:
        country_label = entry.get("country", "")
        zones_map = entry.get("zones") or {}
        zone_letter = zones_map.get(code)
        if not zone_letter or not parse_postal_entry(country_label):
            continue
        label = postal_zone_label(country_label, zone_letter)
        if label and label not in postal_seen[zone_letter]:
            postal_seen[zone_letter].add(label)
            postal_by_zone[zone_letter].append(label)

    zone_letters = set(block.get("rates", {}).get("zones") or [])
    for zone in sorted(zone_letters, key=_zone_letter_sort_key):
        label = f"{direction} {code} Zone {zone}"
        if label not in seen:
            regions.append(label)
            seen.add(label)
        for postal_label in sorted(postal_by_zone.get(zone, [])):
            if postal_label not in seen:
                regions.append(postal_label)
                seen.add(postal_label)

    return regions


def zone_letter_for_region(
    block: dict[str, Any],
    region: str,
    zone_charts: list[dict[str, Any]],
) -> str | None:
    code = service_rate_code(block.get("service") or "")
    direction = direction_label(block)

    zone_match = re.match(
        rf"^{re.escape(direction)} {re.escape(code)} Zone ([A-Z0-9]+)$",
        region,
    )
    if zone_match:
        return zone_match.group(1)

    if is_postal_zone_region(region):
        return zone_letter_from_postal_region(region)

    return None


def build_cost_groups_and_rows(
    blocks: list[dict[str, Any]],
    zone_charts: list[dict[str, Any]],
    home_country: str,
) -> tuple[list[CostColumnGroup], list[MatrixRow]]:
    group_specs: dict[str, dict[str, Any]] = {}
    matrix_rows: dict[str, MatrixRow] = {}

    for block in blocks:
        service = shorten_service_name(block.get("service") or "")
        pricing_model = block.get("pricing_model") or "net"
        currency = block.get("currency") or "EUR"
        entries = block.get("rates", {}).get("entries") or []
        if not entries:
            continue

        entries_by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for entry in entries:
            rate_type = (entry.get("rate_type") or "package").lower()
            entries_by_type[rate_type].append(entry)

        regions = build_regions_for_block(block, zone_charts)

        for rate_type, type_entries in entries_by_type.items():
            cost_key = resolve_cost_key(block, rate_type)
            open_ended = open_ended_weight_labels(type_entries, cost_key=cost_key)
            weight_labels: list[str] = []
            label_prices: dict[str, dict[str, float]] = defaultdict(dict)

            for entry in type_entries:
                weight_label = format_weight_label(
                    entry, pricing_model, open_ended_replacement=open_ended
                )
                if not weight_label:
                    continue
                if weight_label not in weight_labels:
                    weight_labels.append(weight_label)

                zone_prices = entry.get("zones") or {}
                for region in regions:
                    zone_letter = zone_letter_for_region(block, region, zone_charts)
                    if zone_letter is None:
                        continue
                    price = zone_prices.get(zone_letter)
                    if price is None:
                        continue
                    label_prices[region][weight_label] = float(price)

            if not weight_labels:
                continue

            weight_labels.sort(key=weight_sort_key)

            if cost_key not in group_specs:
                group_specs[cost_key] = {
                    "cost_key": cost_key,
                    "currency": currency,
                    "weight_labels": [],
                    "prices": defaultdict(dict),
                    "has_per_kg": False,
                    "has_brackets": False,
                }

            spec = group_specs[cost_key]
            if pricing_model in ("per_pound", "per_kilogram"):
                spec["has_per_kg"] = True
            if any(
                label.lower() not in ("flat", "per kg") for label in weight_labels
            ):
                spec["has_brackets"] = True

            for label in weight_labels:
                if label not in spec["weight_labels"]:
                    spec["weight_labels"].append(label)
            spec["weight_labels"].sort(key=weight_sort_key)

            for region in regions:
                shipment = shipment_fields_for_region(
                    region, block, zone_charts, home_country
                )
                rk = row_key_for(service, shipment)
                if rk not in matrix_rows:
                    matrix_rows[rk] = MatrixRow(
                        origin_country=shipment["origin_country"],
                        origin_country_region=shipment["origin_country_region"],
                        origin_postal_zone=shipment["origin_postal_zone"],
                        destination_country=shipment["destination_country"],
                        destination_country_region=shipment[
                            "destination_country_region"
                        ],
                        destination_postal_zone=shipment["destination_postal_zone"],
                        service=service,
                        currency=currency,
                        row_key=rk,
                    )

                for weight_label in weight_labels:
                    price = label_prices[region].get(weight_label)
                    if price is not None:
                        spec["prices"][rk][weight_label] = price

    cost_groups: list[CostColumnGroup] = []
    for cost_key in COST_KEY_ORDER:
        if cost_key not in group_specs:
            continue
        spec = group_specs[cost_key]
        weight_labels = spec["weight_labels"]
        if spec["has_per_kg"] or spec["has_brackets"]:
            pricing_label = "Rate by Weight/chargeable kg"
        else:
            pricing_label = "Rate by: per shipment"

        cost_groups.append(
            CostColumnGroup(
                cost_key=cost_key,
                pricing_model="net",
                title=transport_cost_title(cost_key),
                pricing_label=pricing_label,
                value_header=value_header_for_group(weight_labels, cost_key),
                currency=spec["currency"],
                weight_labels=weight_labels,
                prices=dict(spec["prices"]),
            )
        )

    rows = sorted(matrix_rows.values(), key=_row_sort_key)
    return cost_groups, rows


class FedExRatesXlsxWriter:
    def __init__(self, data: dict[str, Any]):
        self.data = data
        self.home_country = country_from_source(data)
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Rates"
        self._row = 1
        self._bold = Font(bold=True)
        self._center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self._asymmetric_fill = PatternFill(
            fill_type="solid", fgColor="FF6B6B"
        )

    def write(self) -> Workbook:
        blocks = self.data.get("rate_blocks", [])
        zone_charts = self.data.get("zone_charts", [])
        cost_groups, rows = build_cost_groups_and_rows(
            blocks, zone_charts, self.home_country
        )
        if cost_groups and rows:
            self._write_matrix(cost_groups, rows)
        self._autosize_columns()
        return self.wb

    def _left_col_count(self) -> int:
        return len(SHIPMENT_HEADERS)

    def _write_matrix(
        self,
        cost_groups: list[CostColumnGroup],
        rows: list[MatrixRow],
    ) -> None:
        left_cols = self._left_col_count()
        total_cols = left_cols + sum(cg.width for cg in cost_groups)

        col = left_cols + 1
        for cg in cost_groups:
            self.ws.cell(row=self._row, column=col, value=cg.title)
            self.ws.cell(row=self._row, column=col).font = self._bold
            self.ws.cell(row=self._row, column=col).alignment = self._center
            if cg.width > 1:
                self.ws.merge_cells(
                    start_row=self._row,
                    start_column=col,
                    end_row=self._row,
                    end_column=col + cg.width - 1,
                )
            col += cg.width
        self._row += 1

        col = left_cols + 1
        for cg in cost_groups:
            self.ws.cell(row=self._row, column=col, value=cg.pricing_label)
            self.ws.cell(row=self._row, column=col).font = self._bold
            self.ws.cell(row=self._row, column=col).alignment = self._center
            if cg.width > 1:
                self.ws.merge_cells(
                    start_row=self._row,
                    start_column=col,
                    end_row=self._row,
                    end_column=col + cg.width - 1,
                )
            col += cg.width
        self._row += 1

        col = left_cols + 1
        for cg in cost_groups:
            if len(cg.weight_labels) == 1:
                self.ws.cell(row=self._row, column=col, value=cg.weight_labels[0])
            else:
                for wi, weight in enumerate(cg.weight_labels):
                    cell = self.ws.cell(
                        row=self._row, column=col + 1 + wi, value=weight
                    )
                    cell.font = self._bold
                    cell.alignment = self._center
            col += cg.width
        self._row += 1

        for col_idx, header in enumerate(SHIPMENT_HEADERS, start=1):
            cell = self.ws.cell(row=self._row, column=col_idx, value=header)
            cell.font = self._bold

        col = left_cols + 1
        for cg in cost_groups:
            self.ws.cell(row=self._row, column=col, value="Currency")
            self.ws.cell(row=self._row, column=col).font = self._bold
            value_headers = value_headers_for_labels(cg.weight_labels, cg.cost_key)
            for wi, header in enumerate(value_headers):
                cell = self.ws.cell(
                    row=self._row,
                    column=col + 1 + wi,
                    value=header,
                )
                cell.font = self._bold
            col += cg.width
        self._row += 1

        asymmetric_keys = asymmetric_zone_row_keys(rows)

        for row in rows:
            shipment = [
                row.origin_country,
                row.origin_country_region,
                row.origin_postal_zone,
                row.destination_country,
                row.destination_country_region,
                row.destination_postal_zone,
                row.service,
            ]
            highlight = row.row_key in asymmetric_keys
            for col_idx, value in enumerate(shipment, start=1):
                cell = self.ws.cell(
                    row=self._row, column=col_idx, value=value or None
                )
                if highlight:
                    cell.fill = self._asymmetric_fill

            col = left_cols + 1
            for cg in cost_groups:
                prices = [
                    cg.price_for(row.row_key, weight) for weight in cg.weight_labels
                ]
                has_price = any(price is not None for price in prices)
                currency_cell = self.ws.cell(row=self._row, column=col)
                if has_price:
                    currency_cell.value = row.currency
                if highlight:
                    currency_cell.fill = self._asymmetric_fill
                for wi, weight in enumerate(cg.weight_labels):
                    price = prices[wi]
                    cell = self.ws.cell(
                        row=self._row,
                        column=col + 1 + wi,
                        value=price if price is not None else None,
                    )
                    if isinstance(price, float):
                        cell.number_format = "0.00"
                    if highlight:
                        cell.fill = self._asymmetric_fill
                col += cg.width
            self._row += 1

        _ = total_cols

    def _autosize_columns(self, max_width: int = 24) -> None:
        for col_cells in self.ws.columns:
            letter = get_column_letter(col_cells[0].column)
            length = max(
                (len(str(cell.value)) for cell in col_cells if cell.value is not None),
                default=8,
            )
            self.ws.column_dimensions[letter].width = min(length + 2, max_width)


def data_to_xlsx(data: dict[str, Any], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    writer = FedExRatesXlsxWriter(data)
    wb = writer.write()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def json_to_xlsx(
    json_path: str | Path,
    output_path: str | Path,
) -> Path:
    json_path = Path(json_path)
    with json_path.open(encoding="utf-8") as f:
        data = json.load(f)
    return data_to_xlsx(data, output_path)


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(
        description="Convert FedEx rates JSON to matrix-style XLSX."
    )
    parser.add_argument(
        "input",
        nargs="?",
        default="output/fedex_rates.json",
        help="Input JSON path",
    )
    parser.add_argument(
        "-o",
        "--output",
        default="output/fedex_rates_matrix.xlsx",
        help="Output XLSX path",
    )
    args = parser.parse_args()
    out = json_to_xlsx(args.input, args.output)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
